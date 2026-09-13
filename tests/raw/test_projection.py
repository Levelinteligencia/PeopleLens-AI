"""Testes da camada RAW (F2).

A pergunta que estes testes respondem nao e "o dado esta limpo", e sim: a
sujeira que chegou ao RAW e exatamente a sujeira declarada no catalogo, nem
mais nem menos, e a verdade continua intacta.
"""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import polars as pl
import pytest

from generator import config, materialize, projection, validate_raw
from generator.events.lifecycle import World
from generator.projection import WAVE_1


@pytest.fixture(scope="session")
def projected(tmp_path_factory):
    # A projecao ESCREVE em data/raw. Rodando na raiz do projeto, a suite
    # sobrescrevia o RAW do perfil `dev` por um RAW `smoke` e deixava a camada
    # de verdade intacta: exatamente o achado 1 da F3, fabricado pelos proprios
    # testes. Cada execucao de teste ganha o seu universo descartavel.
    root = tmp_path_factory.mktemp("raw")
    src = Path(config.load().root)
    shutil.copytree(src / "config", root / "config")
    shutil.copytree(src / "data" / "reference", root / "data" / "reference")

    cfg = config.load(root)
    cfg.generation["reproducibility"]["active_profile"] = "smoke"
    tables = materialize.build_all(cfg, World(cfg).run())
    materialize.write(cfg, tables)
    antes = {k: materialize.frame_hash(v) for k, v in tables.items()}
    ledger, manifest = projection.run(cfg, tables)
    depois = {k: materialize.frame_hash(v) for k, v in tables.items()}
    return cfg, tables, ledger, manifest, antes, depois


# ------------------------------------------------------- o RAW nao toca a verdade
def test_projecao_nao_altera_a_camada_de_verdade(projected):
    _, _, _, _, antes, depois = projected
    assert antes == depois


# ------------------------------------------------------------------ disciplina
def test_nenhuma_validacao_bloqueante_do_raw_falha(projected):
    cfg, tables, ledger, _, _, _ = projected
    from generator import validate
    falhas = validate.blocking(validate_raw.run_all(cfg, tables, ledger))
    assert not falhas, "\n".join(f"{c.check_id}: {c.name} | {c.detail}" for c in falhas)


def test_ledger_so_contem_defeitos_do_catalogo(projected):
    cfg, _, ledger, _, _, _ = projected
    conhecidos = {d["id"] for d in cfg.defects["defects"]}
    assert set(ledger.to_frame()["defect_id"].unique().to_list()) <= conhecidos


def test_apenas_sistemas_da_onda_1(projected):
    _, _, ledger, manifest, _, _ = projected
    assert set(manifest["systems"]) == set(WAVE_1)
    assert set(ledger.to_frame()["source_system"].unique().to_list()) <= set(WAVE_1)


# ------------------------------------------------------------------ calibragem
def test_taxa_realizada_proxima_da_esperada(projected):
    _, _, ledger, _, _, _ = projected
    fora = []
    for c in ledger.calibration():
        if c["eligible"] < 500 or c["expected_rate"] <= 0:
            continue
        sd = (c["expected_rate"] * (1 - c["expected_rate"]) / c["eligible"]) ** 0.5
        if abs(c["realized_rate"] - c["expected_rate"]) > max(3 * sd, 0.2 * c["expected_rate"]):
            fora.append(c)
    assert not fora, fora


def test_taxas_ainda_estao_marcadas_como_provisorias(projected):
    """ADR-0008 Parte B: as taxas so podem ser congeladas por decisao da Sam."""
    cfg, _, _, manifest, _, _ = projected
    assert manifest["rates_status"] == "PROVISIONAL"


# ------------------------------------------------------------- formatos nativos
def test_hris_legacy_em_latin1_sem_coluna_de_raca(projected):
    cfg, _, _, _, _, _ = projected
    p = validate_raw._find(cfg.root / "data" / "raw", "HRIS_LEGACY", "employee_master", ".csv")
    assert p is not None
    df = pl.read_csv(p, separator=";", encoding="latin-1", infer_schema_length=0)
    assert df.height > 0
    assert not any(c.lower().startswith(("raca", "race")) for c in df.columns)
    assert set(df["SEXO"].drop_nulls().unique().to_list()) <= {"M", "F", ""}


def test_ats_cloud_produz_json_aninhado(projected):
    cfg, _, _, _, _, _ = projected
    p = validate_raw._find(cfg.root / "data" / "raw", "ATS_CLOUD", "requisition", ".json")
    assert p is not None
    payload = json.loads(p.read_text(encoding="utf-8"))
    assert payload["records"] and "candidates" in payload["records"][0]


def test_vivamarket_nao_expoe_o_id_corporativo(projected):
    cfg, _, _, _, _, _ = projected
    from openpyxl import load_workbook
    p = validate_raw._find(cfg.root / "data" / "raw", "VIVAMARKET_LEGACY", "employee_master", ".xlsx")
    assert p is not None
    ws = load_workbook(p, read_only=True).active
    cols = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    assert "COD_FUNC" in cols and "employee_id" not in cols


def test_colunas_tecnicas_de_lineage_em_todo_arquivo(projected):
    cfg, _, _, manifest, _, _ = projected
    tec = {"_source_file", "_row_number", "_ingested_at", "_source_system"}
    p = validate_raw._find(cfg.root / "data" / "raw", "HRIS_CORE", "employee_master", ".csv")
    df = pl.read_csv(p, separator=",", infer_schema_length=0)
    assert tec <= set(df.columns)


# ------------------------------------------------------------------ fidelidade
def test_defeito_registrado_tem_valor_verdadeiro_e_valor_gravado(projected):
    _, _, ledger, _, _, _ = projected
    led = ledger.to_frame()
    com_campo = led.filter(pl.col("field").is_not_null() & (pl.col("record_key") != "-"))
    assert com_campo.height > 0
    # em um defeito de substituicao, os dois valores existem e sao diferentes
    subst = com_campo.filter(pl.col("defect_id").is_in(["D04", "D12", "D19"]))
    iguais = subst.filter(pl.col("truth_value").eq(pl.col("raw_value"))).height
    assert subst.height > 0 and iguais == 0, f"{iguais} substituicoes que nao mudaram nada"


def test_determinismo_da_projecao(projected):
    cfg, tables, _, _, _, _ = projected
    a, _ = projection.run(cfg, tables)
    b, _ = projection.run(cfg, tables)
    assert materialize.frame_hash(a.to_frame()) == materialize.frame_hash(b.to_frame())


# ------------------------------------------------ relatorio de cobertura (F2)
def test_relatorio_de_cobertura_reconcilia_as_tres_fontes(projected):
    """O relatorio precisa fechar contra catalogo, sources.yaml e ledger."""
    from generator import coverage
    cfg, _, ledger, _, _, _ = projected
    data = coverage.compute(cfg, ledger)
    r = data["resumo"]

    assert r["total"] == len(cfg.defects["defects"])
    assert r["onda_1"] + r["onda_2"] == r["total"]
    assert r["calibracao_fora"] == 0
    assert not r["sem_injecao_na_onda_1"], r["sem_injecao_na_onda_1"]

    # todo sistema em que um defeito foi aplicado tem de constar como declarado
    for d in data["defeitos"]:
        assert set(d["applied_systems"]) <= set(d["declared_systems"]), d["defect_id"]


def test_cobertura_separa_tentativa_de_corrupcao(projected):
    """Achado critico da F2: as duas colunas nao podem ser a mesma coisa."""
    from generator import coverage
    cfg, _, ledger, _, _, _ = projected
    data = coverage.compute(cfg, ledger)
    amostrados = [d for d in data["defeitos"] if d["attempts"] > 0]
    assert amostrados
    for d in amostrados:
        assert d["corrupted"] <= d["attempts"], d["defect_id"]
        assert d["observed_rate"] is not None and d["expected_rate"] is not None


def test_d8_parte_b_segue_provisoria_ate_a_f7(projected):
    cfg, _, _, _, _, _ = projected
    meta = cfg.defects["meta"]
    assert meta["rates_status"] == "PROVISIONAL"
    assert meta["freeze_phase"] == "F7"
