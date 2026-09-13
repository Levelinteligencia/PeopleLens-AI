"""Leitura dos formatos nativos da landing zone.

Princípio 1 e ADR-0003: o RAW nunca e alterado. Este modulo so LE. Ele produz a
camada `staged`, que e uma copia fiel do conteudo, com TUDO como texto, mais um
identificador de linha estavel.

Por que tudo como texto: converter tipo aqui seria interpretar, e interpretar e
trabalho da camada de padronizacao. Um `31/02/2019` precisa chegar inteiro na
proxima camada para ser reprovado com nome e sobrenome, em vez de virar nulo no
carregamento e desaparecer do problema.
"""
from __future__ import annotations

import json
from pathlib import Path

import polars as pl

from generator.config import Config

STAGED = "data/processed/staged"


def _raw_root(cfg: Config) -> Path:
    return Path(cfg.root) / "data" / "raw"


def _find(cfg: Config, system: str, dataset: str, suffix: str) -> Path | None:
    base = _raw_root(cfg) / system / dataset
    if not base.exists():
        return None
    for p in sorted(base.rglob(f"*{suffix}")):
        return p
    return None


def _as_text(df: pl.DataFrame) -> pl.DataFrame:
    return df.with_columns([pl.col(c).cast(pl.Utf8, strict=False) for c in df.columns])


def _with_row_id(df: pl.DataFrame, system: str, dataset: str) -> pl.DataFrame:
    return df.with_columns(
        (pl.lit(f"{system}|{dataset}|") + pl.arange(1, df.height + 1).cast(pl.Utf8)).alias("_row_id")
    )


# --------------------------------------------------------------------------- #
def read_csv(cfg: Config, system: str, dataset: str) -> pl.DataFrame | None:
    meta = cfg.sources["systems"][system]
    p = _find(cfg, system, dataset, ".csv")
    if p is None:
        return None
    df = pl.read_csv(p, separator=meta.get("delimiter", ","),
                     encoding=meta.get("encoding", "utf-8"),
                     infer_schema_length=0, truncate_ragged_lines=True)
    return _with_row_id(_as_text(df), system, dataset)


def read_parquet(cfg: Config, system: str, dataset: str) -> pl.DataFrame | None:
    p = _find(cfg, system, dataset, ".parquet")
    if p is None:
        return None
    return _with_row_id(_as_text(pl.read_parquet(p)), system, dataset)


def read_xlsx(cfg: Config, system: str, dataset: str) -> pl.DataFrame | None:
    from openpyxl import load_workbook

    p = _find(cfg, system, dataset, ".xlsx")
    if p is None:
        return None
    ws = load_workbook(p, read_only=True, data_only=True).active
    it = ws.iter_rows(values_only=True)
    cols = [str(c) for c in next(it)]
    rows = [{c: (None if v is None else str(v)) for c, v in zip(cols, r)} for r in it]
    df = pl.DataFrame(rows, infer_schema_length=None) if rows else pl.DataFrame()
    return _with_row_id(df, system, dataset) if not df.is_empty() else None


def read_json_ats(cfg: Config, system: str) -> dict[str, pl.DataFrame]:
    """O ATS entrega JSON aninhado: requisicao com candidatos dentro.

    O achatamento produz DOIS datasets, respeitando o split do ADR-0005: a
    requisicao tem grao de vaga, a candidatura tem grao de candidato por vaga.
    Achatar tudo num nivel so seria reintroduzir exatamente o erro que aquele
    ADR existe para evitar.
    """
    p = _find(cfg, system, "requisition", ".json")
    if p is None:
        return {}
    payload = json.loads(p.read_text(encoding="utf-8"))
    reqs, apps = [], []
    for r in payload.get("records", []):
        org, dates = r.get("org", {}), r.get("dates", {})
        reqs.append({"requisition_id": r.get("requisition_id"), "position_id": r.get("position_id"),
                     "status": r.get("status"), "recruiter_id": r.get("recruiter_id"),
                     **{f"org_{k}": v for k, v in org.items()},
                     **{k: v for k, v in dates.items()}})
        for c in r.get("candidates", []):
            apps.append({"requisition_id": r.get("requisition_id"), **c})
    out = {}
    if reqs:
        out["requisition"] = _with_row_id(_as_text(pl.DataFrame(reqs, infer_schema_length=None)), system, "requisition")
    if apps:
        out["application"] = _with_row_id(_as_text(pl.DataFrame(apps, infer_schema_length=None)), system, "application")
    return out


# --------------------------------------------------------------------------- #
DATASETS_BY_SYSTEM: dict[str, list[tuple[str, str]]] = {
    "HRIS_LEGACY": [("employee_master", "csv"), ("headcount_snapshot", "csv"), ("movement", "csv")],
    "HRIS_CORE": [("employee_master", "csv"), ("headcount_snapshot", "parquet"),
                  ("performance", "csv"), ("movement", "csv")],
    "PAYROLL_BR": [("payroll_headcount_snapshot", "csv"), ("compensation", "csv")],
    "VIVAMARKET_LEGACY": [("employee_master", "xlsx")],
    "ATS_CLOUD": [("requisition", "json")],
}


def ingest_all(cfg: Config, lineage) -> dict[tuple[str, str], pl.DataFrame]:
    out: dict[tuple[str, str], pl.DataFrame] = {}
    for system, datasets in DATASETS_BY_SYSTEM.items():
        for dataset, fmt in datasets:
            if fmt == "csv":
                df = read_csv(cfg, system, dataset)
                frames = {dataset: df} if df is not None else {}
            elif fmt == "parquet":
                df = read_parquet(cfg, system, dataset)
                frames = {dataset: df} if df is not None else {}
            elif fmt == "xlsx":
                df = read_xlsx(cfg, system, dataset)
                frames = {dataset: df} if df is not None else {}
            else:
                frames = read_json_ats(cfg, system)
            for name, frame in frames.items():
                out[(system, name)] = frame
                lineage.object("dataset", f"staged.{system}.{name}", "file",
                               f"raw.{system}.{name}", f"leitura nativa ({fmt}), sem conversao de tipo",
                               rows=frame.height)
    return out


def write_staged(cfg: Config, staged: dict[tuple[str, str], pl.DataFrame]) -> Path:
    base = Path(cfg.root) / STAGED
    for (system, dataset), df in staged.items():
        d = base / system
        d.mkdir(parents=True, exist_ok=True)
        df.write_parquet(d / f"{dataset}.parquet", compression="zstd")
    return base
